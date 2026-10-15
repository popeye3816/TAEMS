"""
기존 엑셀(유지보수 고객사 현황, 고객사 인수인계 문서)을 읽어 DB 적재용 JSON을 만든다.

사용법:
  python scripts/import/extract.py [--ma 경로] [--handover 경로] [--out .import]

결과 (.import/ 폴더, git에 올라가지 않음):
  name-map.csv   원본 이름 → 고객사 이름 매핑표. Excel로 열어 '고객사(확정)' 열을 고친 뒤 다시 실행한다.
                 '-'를 넣으면 그 원본은 가져오지 않는다.
  extracted.json DB 적재용 데이터 (scripts/import/load.mjs가 읽음)
  검토요약.md     고객사별 건수와 확인이 필요한 항목
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_MA = r"C:\work\doc\유지보수 고객사 현황_202606_v0.3.xlsx"
DEFAULT_HANDOVER = r"C:\work\doc\기술지원_업체정보\고객사인수인계문서_이충호수석.xlsx"

# 자동 매칭이 어려운 이름. (정규화한 원본 이름 → 고객사 이름)
KNOWN_ALIASES = {
    "24년에이치디현대인프라코어웹스피어유지보수계약": "에이치디현대인프라코어",
    "에이치디현대인프라코어": "에이치디현대인프라코어",
    "씨티은행": "한국씨티은행 (인천)",
    "씨티은행모바일개인기업": "한국씨티은행 (인천)",
    "씨티은행콜센터외": "한국씨티은행 (문래동)",
    "씨티은행금융센터외": "한국씨티은행 (문래동)",
    "해군공군정보체계단": "해군/공군사령부",
    "공군해군정보체계단": "해군/공군사령부",
    "조비경농": "조비-경농 (이수시스템)",
    "경농": "조비-경농 (이수시스템)",
    "한화생명금융서비스": "한화생명금융서비스",
    "부산국회도서관": "부산국회도서관",
    "국회부산도서관": "부산국회도서관",
    "케이비펀드파트너스": "KB펀드파트너스",
    "kb펀드파트너스": "KB펀드파트너스",
    "한화손해보험": "한화손해보험",
    "한화손보": "한화손해보험",
    "해군정보체계단": "해군/공군사령부",
    "경농조비이수시스템ehr": "조비-경농 (이수시스템)",
    "경농조비": "조비-경농 (이수시스템)",
    "비엔케이저축은행": "BNK저축은행",
    "그랜드코리아레저": "GKL",
    "처브라이프생명보험": "처브라이프-선릉",
    "신한이지손해보험": "신한EZ손해보험",
    "더케이예다함상조": "TheK예다함-공덕",
    "hyundaiautoevereuropegmbh": "현대오토에버 (유럽법인)",
    "미래에셋대우": "미래에셋증권",
    "미래에셋증권": "미래에셋증권",
    "메트라이프생명보험": "메트라이프생명",
    "한화생명보험": "한화생명",
    "cj푸드빌": "CJ 푸드빌-서울역",
    "대검찰청": "대검찰청 -서초",
}

# 같은 고객사로 추정되지만 확신할 수 없는 이름. 검토요약에 '확인 필요'로 표시된다.
GUESSED_ALIASES = {
    "나이스인프라": "한국전자금융 (Smart ELS)",
    "코리안리재보험": "코리안리재보험",
    "코리안리재보험통합유지보수제니퍼": "코리안리재보험",
    "코리안리": "코리안리재보험",
}

TITLES = (
    "기록연구사|연구사|주무관|매니저|선생님|과장|대리|차장|부장|주임|선임|책임|수석|이사|상무|전무|"
    "사원|팀장|실장|본부장|님"
)
TITLE_RE = re.compile(rf"^\s*([가-힣A-Za-z]{{2,5}})\s*({TITLES})?\s*(.*)$")
PHONE_RE = re.compile(r"0\d{1,2}-\d{3,4}-\d{4}")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
IP_RE = re.compile(r"\d{1,3}(?:\.\d{1,3}){3}")


# ── 공통 유틸 ──────────────────────────────────────────────────────
def s(v) -> str:
    return "" if v is None else str(v).strip()


def lines(v) -> list[str]:
    return [x.strip() for x in s(v).split("\n")]


def norm(name: str) -> str:
    n = name.lower()
    n = re.sub(r"\(주\)|㈜|주식회사", "", n)
    n = re.sub(r"[\s()\[\]/,._\-·]", "", n)
    return n


def norm_core(name: str) -> str:
    """괄호 안 내용까지 지운 핵심 이름 (예: '한국신용평가 (이수시스템)' → '한국신용평가')"""
    return norm(re.sub(r"\(.*?\)", "", name))


def staff_name(v: str) -> str:
    v = re.sub(r"\s*(수석|책임|선임|상무|이사|부장|차장|과장|대리|사원)\s*$", "", s(v))
    return v.strip()


def products_of(*texts: str) -> list[str]:
    t = " ".join(texts).lower()
    out = []
    if "was" in t or "websphere" in t or "웹스피어" in t:
        out.append("WAS")
    if "jennifer" in t or "제니퍼" in t:
        out.append("Jennifer")
    if "j*link" in t:
        out.append("J*Link")
    if "cognos" in t:
        out.append("Cognos")
    return out


def level(v: str) -> str | None:
    v = s(v)
    return v if v in ("상", "중", "하") else None


def remote(v: str) -> str | None:
    v = s(v)
    if not v:
        return None
    if "부분" in v:
        return "부분"
    if "불가" in v:
        return "불가"
    if "가능" in v:
        return "가능"
    return None


def money(v) -> int | None:
    if isinstance(v, (int, float)):
        return int(v)
    d = re.sub(r"[^\d]", "", s(v))
    return int(d) if d else None


def parse_period(text: str) -> tuple[str | None, str | None]:
    parts = re.findall(r"(\d{4})[.\-](\d{1,2})[.\-](\d{1,2})", text)
    def iso(p):
        try:
            return dt.date(int(p[0]), int(p[1]), int(p[2])).isoformat()
        except ValueError:
            return None
    start = iso(parts[0]) if len(parts) >= 1 else None
    end = iso(parts[1]) if len(parts) >= 2 else None
    return start, end


def split_person(line: str) -> dict | None:
    """'홍길동 과장(010-0000-0000)' → {name, title, phone, note}"""
    line = line.strip().strip("-").strip()
    if not line:
        return None
    phone = PHONE_RE.search(line)
    email = EMAIL_RE.search(line)
    core = PHONE_RE.sub("", EMAIL_RE.sub("", line))
    core = re.sub(r"[()]\s*[()]?", " ", core).strip()
    m = TITLE_RE.match(core)
    if not m:
        return None
    name, title, rest = m.group(1), m.group(2), m.group(3).strip(" ()")
    if title == "님":
        title = None
    return {
        "name": name,
        "title": title,
        "phone": phone.group(0) if phone else None,
        "email": email.group(0) if email else None,
        "note": rest or None,
    }


# ── 매핑 ───────────────────────────────────────────────────────────
class NameMap:
    """원본 이름 → 고객사 이름. name-map.csv의 사용자 수정을 우선한다."""

    def __init__(self, path: Path):
        self.path = path
        self.user: dict[tuple[str, str], str] = {}
        self.confirmed: set[tuple[str, str]] = set()
        self.rows: list[dict] = []
        if path.exists():
            with path.open(encoding="utf-8-sig", newline="") as f:
                for r in csv.DictReader(f):
                    self.user[(r["출처"], r["원본이름"])] = r["고객사(확정)"].strip()
                    if "사용자" in r.get("근거", ""):
                        self.confirmed.add((r["출처"], r["원본이름"]))

    def resolve(self, source: str, raw: str, canonical: list[str], hint: str = "", label: str | None = None) -> str | None:
        # label: 매핑표에 표시할 원본 이름 (같은 이름이 여러 건일 때 구분용, 예: 계약번호 포함)
        display = label or raw
        key = (source, display)
        suggestion, reason = self.suggest(raw, canonical)
        # 계약은 EndUser가 본사 이름인 경우가 있어 설명란에 더 구체적인 고객사가 있으면 그쪽을 쓴다
        if hint:
            h = norm(hint)
            hits = [(k, v) for k, v in KNOWN_ALIASES.items() if len(k) >= 5 and k in h]
            hits += [(norm_core(n), n) for n in canonical if len(norm_core(n)) >= 4 and norm_core(n) in h]
            if hits:
                best = max(hits, key=lambda x: len(x[0]))[1]
                if best != suggestion and len(norm_core(best)) > len(norm_core(suggestion)):
                    suggestion, reason = best, "계약 설명 기준 (확인 필요)"
        chosen = self.user.get(key, suggestion)
        if not any(r["출처"] == source and r["원본이름"] == display for r in self.rows):
            self.rows.append({
                "출처": source,
                "원본이름": display,
                "고객사(확정)": chosen,
                "자동제안": suggestion,
                "근거": "사용자 확인" if key in self.confirmed
                        else reason if key not in self.user
                        else ("사용자 수정" if chosen != suggestion else reason),
            })
        return None if chosen in ("", "-") else chosen

    @staticmethod
    def suggest(raw: str, canonical: list[str]) -> tuple[str, str]:
        n, c = norm(raw), norm_core(raw)
        if n in KNOWN_ALIASES:
            return KNOWN_ALIASES[n], "알려진 별칭"
        if c in KNOWN_ALIASES:
            return KNOWN_ALIASES[c], "알려진 별칭"
        if n in GUESSED_ALIASES or c in GUESSED_ALIASES:
            return GUESSED_ALIASES.get(n) or GUESSED_ALIASES[c], "추정 (확인 필요)"
        for name in canonical:
            if norm(name) == n:
                return name, "이름 일치"
        for name in canonical:
            if norm_core(name) == c and c:
                return name, "괄호 제외 일치"
        cands = [name for name in canonical if c and (c in norm_core(name) or norm_core(name) in c) and len(norm_core(name)) >= 3]
        if len(cands) == 1:
            return cands[0], "부분 일치 (확인 필요)"
        clean = re.sub(r"\(주\)|㈜|주식회사", "", raw).strip()
        return clean, "새 고객사 (확인 필요)"

    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.path.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["출처", "원본이름", "고객사(확정)", "자동제안", "근거"])
            w.writeheader()
            w.writerows(self.rows)


# ── 원본 읽기 ──────────────────────────────────────────────────────
def read_ma(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    ma_rows, contracts = [], []

    ws = wb["MA"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    names = [s(r[5]) for r in rows if s(r[5])]
    dup = {n for n in names if names.count(n) > 1}
    for r in rows:
        name = s(r[5])
        if not name:
            continue
        business = s(r[6])
        display = f"{name} ({business})" if name in dup and business else name
        ma_rows.append({
            "raw_name": display,
            "group_name": name if name in dup else None,
            "business_unit": s(r[0]) or None,
            "sales_rep": staff_name(r[1]) or None,   # 2번째 열만 사용 (16번째 열은 사용하지 않음)
            "importance": level(r[2]),
            "annual_amount": money(r[3]) or money(r[4]),
            "business_scope": business or None,
            "product_text": s(r[7]),
            "version_text": s(r[8]),
            "cycle": s(r[9]) or None,
            "remote": remote(r[10]),
            "support_level": "2차" if s(r[12]) else ("1차" if s(r[11]) else None),
            "support_count": s(r[13]) or None,
            "difficulty": level(r[14]),
            "engineer": staff_name(r[16]) or None,
            "contact_text": s(r[17]),
            "contact_phone": s(r[18]),
            "notes": s(r[19]).lstrip("- ").strip() or None,
        })

    ws = wb["계약"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        no = s(r[0])
        if not no.startswith("PJT"):
            continue
        period = s(r[1])
        start, end = parse_period(period)
        desc = s(r[4])
        contracts.append({
            "contract_no": no,
            "period_text": period or None,
            "start_date": start,
            "end_date": end,
            "contractor": s(r[2]) or None,
            "end_user": s(r[3]) or None,
            "description": desc or None,
            "amount": money(r[5]),
            "billing_note": s(r[9]) or None,
            "sales_team": s(r[13]) or None,
            "sales_rep": staff_name(r[14]) or None,
            "engineer": staff_name(r[26]) if len(r) > 26 else None,
            "products": s(r[27]) if len(r) > 27 else None,
            "auto_renew": "자동연장" in desc,
        })
    return ma_rows, contracts


SHEET_FIELDS = {
    "고객사 위치": "address",
    "담당자 성향": "contact_temperament",
    "지원 빈도": "support_freq",
    "지원 방법": "support_how",
    "사이트 방문 시 출입 절차": "visit_procedure",
    "대상 시스템 환경": "system_env_note",
    "계약 제품 외 추가 지원 요건": "extra_requirements",
    "정기점검 방법": "inspection_method",
    "장애 발생시 조치 절차": "incident_procedure",
    "고객 담당자 성명": "contacts_text",
    "담당영업": "sales_rep",
    "계약 종류": "contract_type",
}
SHEET_NOTES = [
    "사이트 구성 / 설정 특이사항", "최근 3개월 구성/설정 변경 내역", "최근 3개월 장애 발생 유형",
    "미해결 이슈", "기타 사이트 특이사항", "기타 사이트  정보", "기타 사이트 정보", "기타 참고사항",
]
EMPTY = {"", "없음", "-", "x", "X"}


def read_handover(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    handover_date = dt.date.fromtimestamp(os.path.getmtime(path)).isoformat()

    summary = []
    for r in wb["요약"].iter_rows(min_row=5, values_only=True):
        name = s(r[2])
        if not name or not s(r[1]).isdigit():
            continue
        summary.append({
            "raw_name": name,
            "col_system": lines(r[3]),
            "col_ip": lines(r[4]),
            "col_version": lines(r[5]),
            "contact_names": lines(r[6]),
            "contact_phones": lines(r[7]),
            "contact_emails": lines(r[8]),
            "address": s(r[9]) or None,
            "sales_rep": staff_name(r[10]) or None,
            "kind": s(r[11]),
            "cycle": s(r[12]) or None,
            "remark": s(r[13]) or None,
        })

    sheets = []
    for title in wb.sheetnames[1:]:
        ws = wb[title]
        info = {"sheet": title, "fields": {}, "notes": []}
        for r in ws.iter_rows(values_only=True):
            label, value = s(r[1]), s(r[2]) if len(r) > 2 else ""
            if s(r[2]) == "인계자 :" or label == "인계자 :":
                pass
            if len(r) > 3 and s(r[2]) in ("인계자 :", "인수자 :"):
                info["fields"]["from" if "인계" in s(r[2]) else "to"] = staff_name(r[3])
                continue
            if label == "고객사 이름":
                info["name"] = value
                continue
            flat = label.replace("\n", " ")
            for key, field in SHEET_FIELDS.items():
                if flat.startswith(key) and value not in EMPTY:
                    info["fields"][field] = value
            for note_label in SHEET_NOTES:
                if flat.strip() == note_label.strip() and value not in EMPTY:
                    info["notes"].append(f"[{note_label.strip()}] {value}")
        if info.get("name"):
            sheets.append(info)
    return summary, sheets, handover_date


# ── 변환 ───────────────────────────────────────────────────────────
def contacts_from_columns(names, phones, emails) -> list[dict]:
    out, org = [], None
    for i, line in enumerate(names):
        phone = phones[i] if i < len(phones) else ""
        email = emails[i] if i < len(emails) else ""
        if not line:
            continue
        # '통합유지보수담당(우리IT)'처럼 연락처 없는 소속 줄은 다음 사람들의 소속으로 쓴다
        if not phone and not email and ("담당" in line or "업체" in line) and not TITLE_RE.match(line).group(2):
            org = line
            continue
        for part_idx, part in enumerate(re.split(r"\s*->\s*", line)):
            p = split_person(part)
            if not p:
                continue
            is_old = "->" in line and part_idx == 0
            if not is_old:
                p["phone"] = p["phone"] or (PHONE_RE.search(phone).group(0) if PHONE_RE.search(phone) else (phone or None))
                p["email"] = p["email"] or (EMAIL_RE.search(email).group(0) if EMAIL_RE.search(email) else None)
            sub_org = re.search(r"\(([^)]*(업체|웰클라우드|isecu)[^)]*)\)", line)
            p["organization"] = sub_org.group(1) if sub_org else org
            p["is_active"] = not is_old
            out.append(p)
    return out


def contacts_from_text(text: str) -> list[dict]:
    out, org = [], None
    for line in lines(text):
        if not line:
            continue
        if not PHONE_RE.search(line) and ("담당" in line or "업체" in line):
            org = line
            continue
        p = split_person(line)
        if p:
            p["organization"] = org
            p["is_active"] = True
            out.append(p)
    return out


def systems_from_summary(h: dict, products: list[str]) -> list[dict]:
    product = products[0] if len(products) == 1 else ("Jennifer" if "Jennifer" in products else (products[0] if products else "기타"))
    sys_col, ip_col, ver_col = h["col_system"], h["col_ip"], h["col_version"]
    # IP가 '시스템' 열에 들어간 행도 있다
    if sys_col and all(IP_RE.search(x) or x.startswith("http") for x in sys_col if x):
        ip_col, sys_col = sys_col, []
    out = []
    count = max(len(ip_col), len(sys_col) if any(IP_RE.search(x) for x in ip_col) else 0)
    for i in range(count):
        ip = ip_col[i] if i < len(ip_col) else ""
        label = sys_col[i] if i < len(sys_col) else ""
        if not IP_RE.search(ip):
            continue
        version = ver_col[i] if i < len(ver_col) else (ver_col[0] if len(ver_col) == 1 else "")
        out.append({
            "business_name": label or None,
            "ip": ip,
            "product": product,
            "version": version or None,
            "note": None,
        })
    return out


def systems_from_ma(row: dict) -> list[dict]:
    out = []
    for p in products_of(row["product_text"], row["version_text"]):
        out.append({"business_name": row["business_scope"], "ip": None, "product": p,
                    "version": row["version_text"] or None, "note": "엑셀 MA 시트 기준 (서버 정보 미입력)"})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ma", default=DEFAULT_MA)
    ap.add_argument("--handover", default=DEFAULT_HANDOVER)
    ap.add_argument("--out", default=".import")
    args = ap.parse_args()

    out_dir = Path(args.out)
    ma_rows, contracts = read_ma(args.ma)
    summary, sheets, handover_date = read_handover(args.handover)
    nm = NameMap(out_dir / "name-map.csv")

    canonical = [r["raw_name"] for r in ma_rows]
    customers: dict[str, dict] = {}
    warnings: list[str] = []

    def customer(name: str) -> dict:
        if name not in customers:
            customers[name] = {"name": name, "products": [], "sources": [], "contacts": [], "systems": [],
                               "assignments": [], "site_notes": []}
        return customers[name]

    # 1) MA 시트
    for r in ma_rows:
        name = nm.resolve("MA", r["raw_name"], canonical)
        if not name:
            continue
        c = customer(name)
        merged = "MA" in c["sources"]  # MA의 여러 행을 한 고객사로 합치는 경우
        c["sources"].append("MA")
        for k in ("group_name", "business_unit", "sales_rep", "importance", "annual_amount",
                  "business_scope", "support_level"):
            if r[k] is not None and c.get(k) is None:
                c[k] = r[k]
        def join(field, value):
            if value and value not in (c.get(field) or ""):
                c[field] = " / ".join(x for x in (c.get(field), value) if x)
        if not merged:
            c["remote_access"] = r["remote"]
            c["service_difficulty"] = r["difficulty"]
        join("notes", r["notes"])
        join("product_versions", " / ".join(x for x in (r["product_text"], r["version_text"]) if x) or None)
        join("inspection_cycle_note", r["cycle"])
        join("support_count_note", r["support_count"])
        c["products"] = sorted(set(c["products"]) | set(products_of(r["product_text"], r["version_text"])))
        c["_ma_row"] = r
        if r["engineer"] and not any(a["engineer"] == r["engineer"] and a["ended_on"] is None for a in c["assignments"]):
            c["assignments"].append({"engineer": r["engineer"], "role": "primary", "started_on": "2026-06-01",
                                     "ended_on": None, "note": "유지보수 고객사 현황(2026-06) 기준"})
        if r["contact_text"] and r["contact_text"] != "모름":
            for p in contacts_from_text(r["contact_text"]):
                if r["contact_phone"] and not p["phone"]:
                    p["phone"] = r["contact_phone"]
                c["contacts"].append(p)

    # 2) 인수인계 요약
    canonical = list(customers)
    for h in summary:
        name = nm.resolve("인수인계 요약", h["raw_name"], canonical)
        if not name:
            continue
        c = customer(name)
        c["sources"].append("인수인계")
        prods = products_of(h["kind"])
        c["products"] = sorted(set(c["products"]) | set(prods))
        c["address"] = h["address"] or c.get("address")
        c["sales_rep"] = c.get("sales_rep") or h["sales_rep"]
        if h["cycle"] and h["cycle"] != c.get("inspection_cycle_note"):
            c["inspection_cycle_note"] = " / ".join(x for x in (c.get("inspection_cycle_note"), f"인수인계: {h['cycle']}") if x)
        if h["remark"] and "계약종료" in h["remark"]:
            c["status"] = "ended"
        ho_contacts = contacts_from_columns(h["contact_names"], h["contact_phones"], h["contact_emails"])
        if ho_contacts:
            # 인수인계 문서가 더 최신이므로 MA 담당자를 대체한다
            c["contacts"] = ho_contacts
        systems = systems_from_summary(h, prods)
        if systems:
            c["systems"] = systems
        h_from = re.search(r"([가-힣]{2,4})사원\s*인수인계", h["remark"] or "")
        c["_handover_from"] = h_from.group(1) if h_from else None
        if any(not x for x in h["contact_phones"]) and len(h["contact_names"]) > len([x for x in h["contact_phones"] if x]) + 1:
            warnings.append(f"{name}: 담당자 이름과 연락처 줄 수가 달라 연결이 어긋났을 수 있음")

    # 3) 인수인계 개별 시트
    canonical = list(customers)
    for sh in sheets:
        name = nm.resolve("인수인계 시트", sh["name"], canonical)
        if not name:
            continue
        c = customer(name)
        if "인수인계" not in c["sources"]:
            c["sources"].append("인수인계")
        f = sh["fields"]
        for field in ("address", "contact_temperament", "visit_procedure", "system_env_note", "extra_requirements",
                      "inspection_method", "incident_procedure"):
            if f.get(field):
                c[field] = f[field]
        method = [x for x in (f.get("support_freq"), f.get("support_how")) if x]
        if method:
            c["support_method"] = "\n".join(method)
        if f.get("contract_type"):
            c["site_notes"].append(f"[계약 종류] {f['contract_type']}")
        c["site_notes"].extend(sh["notes"])
        if not c.get("sales_rep") and f.get("sales_rep"):
            c["sales_rep"] = staff_name(f["sales_rep"])
        if not c["contacts"] and f.get("contacts_text"):
            c["contacts"] = contacts_from_text(f["contacts_text"])
        to_eng = f.get("to")
        from_eng = c.pop("_handover_from", None) or f.get("from")
        if to_eng:
            prev = [a for a in c["assignments"] if a["role"] == "primary" and a["ended_on"] is None]
            for a in prev:
                if a["engineer"] != to_eng:
                    a["ended_on"] = handover_date
                    a["note"] = (a["note"] or "") + f" → {to_eng}에게 인수인계"
            if not prev and from_eng and from_eng != to_eng:
                c["assignments"].append({"engineer": from_eng, "role": "primary", "started_on": "2026-01-01",
                                         "ended_on": handover_date, "note": f"인수인계서 인계자 → {to_eng}"})
            if not any(a["engineer"] == to_eng and a["ended_on"] is None for a in c["assignments"]):
                c["assignments"].append({"engineer": to_eng, "role": "primary", "started_on": handover_date,
                                         "ended_on": None, "note": f"인수인계서 ({from_eng or '인계자 미상'} → {to_eng})"})

    # MA에만 있는 고객사는 제품 단위 시스템 행을 만든다
    for c in customers.values():
        if not c["systems"] and c.get("_ma_row"):
            c["systems"] = systems_from_ma(c["_ma_row"])
        c.pop("_ma_row", None)
        c.pop("_handover_from", None)
        c["site_notes"] = "\n".join(c["site_notes"]) or None
        if c["contacts"]:
            actives = [p for p in c["contacts"] if p["is_active"]]
            if actives:
                actives[0]["is_primary"] = True

    # 4) 계약
    canonical = list(customers)
    contract_out = []
    for k in contracts:
        name = nm.resolve("계약", k["end_user"] or "", canonical, hint=k["description"] or "",
                           label=f"{k['end_user']} ({k['contract_no']})") if k["end_user"] else None
        if name:
            c = customer(name)
            if "계약" not in c["sources"]:
                c["sources"].append("계약")
            c["products"] = sorted(set(c["products"]) | set(products_of(k["products"] or "")))
            c["sales_rep"] = c.get("sales_rep") or k["sales_rep"]
            c["business_unit"] = c.get("business_unit") or k["sales_team"]
            if not c["systems"]:
                c["systems"] = [{"business_name": None, "ip": None, "product": p, "version": None,
                                 "note": "계약 정보 기준 (서버 정보 미입력)"} for p in products_of(k["products"] or "")]
        contract_out.append({**k, "customer": name})

    # 5) 직원 명단
    engineers = sorted({a["engineer"] for c in customers.values() for a in c["assignments"]} |
                       {k["engineer"] for k in contracts if k["engineer"]})
    sales = sorted({c["sales_rep"] for c in customers.values() if c.get("sales_rep")} |
                   {k["sales_rep"] for k in contracts if k["sales_rep"]})

    data = {
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "handover_date": handover_date,
        "staff": [{"name": n, "kind": "engineer"} for n in engineers] + [{"name": n, "kind": "sales"} for n in sales],
        "customers": sorted(customers.values(), key=lambda c: c["name"]),
        "contracts": contract_out,
    }
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "extracted.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    nm.save()

    # 검토 요약
    need = [r for r in nm.rows if "확인 필요" in r["근거"]]
    md = [f"# 엑셀 가져오기 검토 요약 ({data['generated_at']})", "",
          f"- 고객사 {len(customers)}곳, 계약 {len(contract_out)}건 (고객사 미연결 {sum(1 for k in contract_out if not k['customer'])}건)",
          f"- 엔지니어 {len(engineers)}명: {', '.join(engineers)}",
          f"- 영업 {len(sales)}명: {', '.join(sales)}",
          f"- 담당자 {sum(len(c['contacts']) for c in customers.values())}명, "
          f"시스템 {sum(len(c['systems']) for c in customers.values())}건 "
          f"(IP 있음 {sum(1 for c in customers.values() for x in c['systems'] if x['ip'])}건)",
          "", "## 이름 매핑 확인 필요", "", "| 출처 | 원본 이름 | 제안 고객사 | 근거 |", "|---|---|---|---|"]
    md += [f"| {r['출처']} | {r['원본이름']} | {r['고객사(확정)']} | {r['근거']} |" for r in need]
    md += ["", "## 경고", ""] + [f"- {w}" for w in warnings]
    md += ["", "## 고객사별 건수", "", "| 고객사 | 출처 | 제품 | 영업 | 현재 담당 | 담당자 | 시스템 |", "|---|---|---|---|---|---|---|"]
    for c in data["customers"]:
        cur = ", ".join(a["engineer"] for a in c["assignments"] if a["ended_on"] is None)
        md.append(f"| {c['name']} | {'+'.join(c['sources'])} | {', '.join(c['products'])} | {c.get('sales_rep') or ''} "
                  f"| {cur} | {len(c['contacts'])} | {len(c['systems'])} |")
    (out_dir / "검토요약.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    print(md[2]); print(md[3]); print(md[4]); print(md[5])
    print(f"- 이름 매핑 확인 필요: {len(need)}건, 경고: {len(warnings)}건")
    print(f"결과: {out_dir / 'name-map.csv'}, {out_dir / 'extracted.json'}, {out_dir / '검토요약.md'}")


if __name__ == "__main__":
    main()
